"""Generate the Milestone 3 slide deck: Milestone_3_Resume_Matcher.pptx (project root).

Usage (from the project root):
    .venv/bin/pip install python-pptx      # one-time; not needed by the app itself
    .venv/bin/python presentation/build_presentation.py

The slides use simple language. Every number comes from the repository: dataset sizes from the data
loader, Precision@5 and top-1 results read from results/evaluation_summary.csv (`python evaluate.py`),
and the test count from `python -m pytest` on 27 Oct 2026. The app screenshot in
presentation/assets/app_screenshot.png is a real capture of the running Streamlit app (demo_samples/ input).
"""

from __future__ import annotations

import csv
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
ASSETS = Path(__file__).resolve().parent / "assets"
OUTPUT = ROOT / "Milestone_3_Resume_Matcher.pptx"
NOTES_OUTPUT = Path(__file__).resolve().parent / "speaker_notes.md"
RESULTS_CSV = ROOT / "results" / "evaluation_summary.csv"
REPO_URL = "https://github.com/soniraghav217-cmyk/Resume-Matcher.git"
TOTAL_SLIDES = 7
TEST_COUNT = 119  # `python -m pytest`, 27 Oct 2026

# Palette: deep navy ground, white text, cyan as the one sharp accent, blue/violet/green for the methods.
NAVY = RGBColor(0x0A, 0x12, 0x26)
PANEL = RGBColor(0x11, 0x1F, 0x3A)
PANEL_2 = RGBColor(0x16, 0x29, 0x4B)
BORDER = RGBColor(0x24, 0x3B, 0x63)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ICE = RGBColor(0xC9, 0xDB, 0xF5)
MUTED = RGBColor(0x8F, 0xA3, 0xC2)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
BLUE = RGBColor(0x3B, 0x82, 0xF6)
VIOLET = RGBColor(0xA7, 0x8B, 0xFA)
GREEN = RGBColor(0x34, 0xD3, 0x99)
METHOD_KEYS = ("tfidf", "glove", "sbert", "hybrid")
METHOD_NAMES = ["TF-IDF", "GloVe", "Sentence-BERT", "Hybrid"]
METHOD_COLORS = (CYAN, BLUE, VIOLET, GREEN)

FONT = "Calibri"
MONO = "Courier New"

# Speaker notes, one per slide in order. Written to speaker_notes.md instead of the .pptx because
# Keynote refuses to import decks that contain python-pptx notes slides ("file format is invalid").
NOTES: list[str] = []


# --------------------------------------------------------------------------- assets


def make_background(path: Path, glows: list[tuple[int, int, int, tuple[int, int, int]]]) -> None:
    """Navy vertical gradient + soft blurred glows + faint dot grid (1920x1080)."""
    w, h = 1920, 1080
    t = np.linspace(0, 1, h)[:, None, None]
    top, bottom = np.array([9, 16, 34]), np.array([5, 9, 20])
    base = (top + (bottom - top) * t) * np.ones((h, w, 3))
    glow = Image.new("RGB", (w, h))
    draw = ImageDraw.Draw(glow)
    for cx, cy, r, color in glows:
        draw.ellipse((cx - r, cy - r, cx + r, cy + r), fill=color)
    glow = np.asarray(glow.filter(ImageFilter.GaussianBlur(220)), dtype=float)
    image = Image.fromarray(np.clip(base + glow, 0, 255).astype("uint8"))
    dots = ImageDraw.Draw(image)
    for y in range(30, h, 48):
        for x in range(30, w, 48):
            dots.point((x, y), fill=(28, 44, 72))
    image.save(path, optimize=True)


def ensure_assets() -> None:
    ASSETS.mkdir(exist_ok=True)
    make_background(ASSETS / "bg_title.png", [(1500, 330, 420, (10, 70, 110)), (300, 980, 380, (20, 40, 110))])
    make_background(ASSETS / "bg_content.png", [(1750, 80, 380, (8, 48, 82)), (120, 1050, 300, (14, 28, 80))])
    if not (ASSETS / "app_screenshot.png").is_file():
        raise SystemExit("Missing presentation/assets/app_screenshot.png (real app screenshot).")


# --------------------------------------------------------------------------- helpers


def text(slide, x, y, w, h, content, size=16, color=WHITE, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, italic=False, name=None):
    """Add a text box.

    ``content`` is a string, or a list of paragraphs where each paragraph is a string or a
    sequence of ``(text, overrides)`` runs (overrides: size, bold, italic, color, font).
    """
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paragraphs = content if isinstance(content, list) else [content]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = [(para, {})] if isinstance(para, str) else para
        for run_text, opts in runs:
            r = p.add_run()
            r.text = run_text
            f = r.font
            f.name = opts.get("font", font)
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.italic = opts.get("italic", italic)
            f.color.rgb = opts.get("color", color)
    return box


def bullets(slide, x, y, w, h, items, size=15, color=WHITE, bullet_color=CYAN, space_after=6):
    """Bulleted list using real paragraph bullets (no literal bullet characters in the text)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.24))))
        pPr.set("indent", str(-Emu(Inches(0.24))))
        bu_clr = pPr.makeelement(qn("a:buClr"), {})
        clr = bu_clr.makeelement(qn("a:srgbClr"), {"val": str(bullet_color)})
        bu_clr.append(clr)
        pPr.append(bu_clr)
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
        r = p.add_run()
        r.text = item
        r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(size), color
    return box


def card(slide, x, y, w, h, fill=PANEL, border=BORDER, radius=0.08, name=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if border is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = border
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    if name:
        shape.name = name
    return shape


def badge(slide, x, y, d, label, fill=CYAN, color=NAVY, size=16):
    """Filled circle with a short label (number or symbol) — the deck's repeated motif."""
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), True, color
    return c


def pill(slide, x, y, w, h, label, fill=PANEL_2, color=ICE, border=BORDER, size=13, bold=False):
    shape = card(slide, x, y, w, h, fill=fill, border=border, radius=0.5)
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, color
    return shape


def arrow(slide, x, y, w=0.26, h=0.3, color=CYAN):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()
    a.shadow.inherit = False
    return a


def bar(slide, x, y, w, h, color):
    b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.adjustments[0] = 0.5
    b.fill.solid()
    b.fill.fore_color.rgb = color
    b.line.fill.background()
    b.shadow.inherit = False
    return b


def content_slide(prs, number, eyebrow, title):
    """Shared frame for slides 2-7: background, eyebrow, title, footer and slide number."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(ASSETS / "bg_content.png"), 0, 0, prs.slide_width, prs.slide_height)
    text(slide, 0.6, 0.38, 9, 0.3, eyebrow.upper(), size=13, color=CYAN, bold=True, name="Eyebrow")
    text(slide, 0.6, 0.66, 12.1, 0.7, title, size=34, bold=True, name="Title")
    text(slide, 0.6, 7.0, 7, 0.3, "Resume Matcher  ·  Milestone 3  ·  Batch 3", size=11, color=MUTED, name="Footer")
    text(slide, 11.7, 7.0, 1.0, 0.3, f"{number} / {TOTAL_SLIDES}", size=11, color=MUTED, align=PP_ALIGN.RIGHT,
         name="Slide number")
    return slide



# --------------------------------------------------------------------------- slides


def load_results() -> dict[tuple[str, str], dict[str, str]]:
    """Rows of results/evaluation_summary.csv keyed by (dataset, method)."""
    with open(RESULTS_CSV, newline="", encoding="utf-8") as f:
        return {(row["dataset"], row["method"]): row for row in csv.DictReader(f)}


def slide_title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(ASSETS / "bg_title.png"), 0, 0, prs.slide_width, prs.slide_height)

    text(slide, 0.7, 1.15, 6.6, 0.35, "B.TECH CSE  ·  MILESTONE 3 EVALUATION", size=14, color=CYAN, bold=True)
    text(slide, 0.7, 1.55, 6.8, 1.2, "Resume Matcher", size=66, bold=True)
    text(slide, 0.7, 2.85, 6.3, 0.9, "Finds the best resumes for a job, automatically, by comparing their text",
         size=22, color=ICE)
    widths = (1.15, 1.05, 1.75, 1.1)
    for i, (label, color) in enumerate(zip(METHOD_NAMES, METHOD_COLORS)):
        x = 0.7 + sum(widths[:i]) + 0.15 * i
        pill(slide, x, 3.95, widths[i], 0.42, label, fill=PANEL, color=color, border=color, size=13, bold=True)

    text(slide, 0.7, 4.75, 6, 0.3, "TEAM  ·  BATCH NUMBER 3", size=12, color=MUTED, bold=True)
    for i, (name, enrol) in enumerate((("Yash Agarwal", "S24CSEU0244"), ("Raghav Soni", "S24CSEU0303"))):
        x = 0.7 + i * 3.25
        card(slide, x, 5.1, 3.05, 1.05, fill=PANEL)
        badge(slide, x + 0.2, 5.37, 0.5, name.split()[0][0] + name.split()[1][0], fill=CYAN, size=14)
        text(slide, x + 0.85, 5.27, 2.1, 0.4, name, size=18, bold=True)
        text(slide, x + 0.85, 5.67, 2.1, 0.35, enrol, size=14, color=CYAN, font=MONO, bold=True)
    text(slide, 0.7, 6.55, 6, 0.3, "27 October 2026", size=13, color=MUTED)

    # Illustrative visual: a job description being matched to ranked resumes.
    card(slide, 7.75, 1.9, 1.9, 2.75, fill=PANEL, border=CYAN, name="Visual: job description")
    text(slide, 7.95, 2.05, 1.6, 0.3, "Job Description", size=13, bold=True, color=CYAN)
    for i, frac in enumerate((0.95, 0.8, 0.9, 0.6, 0.85, 0.7, 0.5)):
        bar(slide, 7.95, 2.5 + i * 0.27, 1.5 * frac, 0.09, BORDER if i % 3 else MUTED)
    badge(slide, 9.95, 2.82, 0.9, "AI", fill=PANEL_2, color=CYAN, size=20)
    arrow(slide, 9.7, 3.12, 0.24, 0.3)
    arrow(slide, 10.9, 3.12, 0.24, 0.3)
    for i, (frac, color) in enumerate(((0.92, CYAN), (0.72, BLUE), (0.52, VIOLET))):
        y = 1.55 + i * 1.18
        card(slide, 11.25, y, 1.55, 1.0, fill=PANEL, name=f"Visual: ranked resume {i + 1}")
        badge(slide, 11.38, y + 0.12, 0.42, f"#{i + 1}", fill=color, size=12)
        bar(slide, 11.9, y + 0.2, 0.75, 0.08, MUTED)
        bar(slide, 11.9, y + 0.38, 0.55, 0.08, BORDER)
        bar(slide, 11.38, y + 0.72, 1.3, 0.1, BORDER)
        bar(slide, 11.38, y + 0.72, 1.3 * frac, 0.1, color)
    text(slide, 7.75, 5.15, 5.05, 0.6, "You give it a job and some resumes. It sorts the resumes from best to worst match.",
         size=13, color=MUTED, italic=True, align=PP_ALIGN.CENTER)

    NOTES.append(
        "Good morning. We are Batch 3: Yash Agarwal and Raghav Soni. Our project is Resume Matcher. "
        "In one line: you give it a job description and some resumes, and it sorts the resumes from the best "
        "match to the worst match. It does this by comparing the text of the job with the text of each resume. "
        "We built and compared four ways of doing this comparison: TF-IDF, GloVe, Sentence-BERT, and a Hybrid "
        "that combines all three."
    )


def slide_problem(prs):
    slide = content_slide(prs, 2, "The Problem · Our Idea · The Data", "Why do we need this?")

    for i, (label, title, body) in enumerate((
        ("?", "The problem",
         "A company can get hundreds of resumes for one job. Reading all of them takes a lot of time. "
         "A simple keyword search also misses good people who write the same skill in other words "
         "(for example \"ML\" instead of \"Machine Learning\")."),
        ("→", "Our idea",
         "Give the app a job description and some resumes. It gives every resume a score from 0 to 100 "
         "and sorts them from best to worst match. The recruiter starts with the top ones. "
         "A person still makes the final decision."),
    )):
        x = 0.6 + i * 6.225
        card(slide, x, 1.5, 5.875, 2.05)
        badge(slide, x + 0.25, 1.75, 0.55, label, size=20)
        text(slide, x + 1.0, 1.68, 4.65, 0.4, title, size=20, bold=True)
        text(slide, x + 1.0, 2.12, 4.65, 1.35, body, size=15, color=ICE)

    text(slide, 0.6, 3.82, 12, 0.3, "THE DATA WE USED TO TEST IT  (real, public resume datasets from Kaggle)",
         size=12, color=MUTED, bold=True)
    stats = (
        ("166", "different tech resumes", "Tech jobs dataset  (our main one)",
         "The file has 962 rows, but 796 were exact copies, so 166 are unique. "
         "25 job types: Java / Python / .NET Developer, Data Science, DevOps, Testing, Database ..."),
        ("2,481", "resumes, all industries", "All-industries dataset  (extra check)",
         "2,484 rows; we removed 1 empty and 2 copies. 24 broad areas such as IT, Engineering, HR, "
         "Finance. Very few of them are software resumes."),
    )
    for i, (big, unit, title, detail) in enumerate(stats):
        x = 0.6 + i * 6.225
        card(slide, x, 4.18, 5.875, 1.75, fill=PANEL)
        text(slide, x + 0.3, 4.28, 2.2, 0.85, big, size=46, bold=True, color=CYAN if i == 0 else BLUE)
        text(slide, x + 0.3, 5.13, 2.2, 0.6, unit, size=13, color=MUTED)
        text(slide, x + 2.55, 4.38, 3.1, 0.35, title, size=15, bold=True)
        text(slide, x + 2.55, 4.78, 3.1, 1.1, detail, size=12, color=ICE)

    text(slide, 0.6, 6.15, 12.1, 0.6,
         [(("Why two datasets? ", {"bold": True, "color": CYAN}),
           ("Our project is about tech jobs (software, data, AI, testing), so the tech dataset is the main one. "
            "The bigger mixed dataset shows how the methods behave when most resumes are not from tech.", {}))],
         size=13, color=ICE)

    NOTES.append(
        "First, the problem. For one job opening, a company can get hundreds of resumes. Reading every one by hand "
        "is slow and tiring. Searching for keywords is faster, but it misses people who write the same skill in "
        "different words, like 'ML' instead of 'Machine Learning'.\n\n"
        "Our idea: the recruiter gives our app the job description and the resumes. The app gives each resume a "
        "score from 0 to 100 and sorts them. The recruiter can start reading from the top. The app only helps; "
        "a human still makes the final decision.\n\n"
        "To test it we used two real, public resume datasets from Kaggle. The tech dataset is our main one. It "
        "has 962 rows, but 796 of them were exact copies, so after removing copies we have 166 unique resumes in "
        "25 job types. The second dataset has about 2,500 resumes from all industries. We use it as an extra "
        "check. Our code counts and cleans these numbers itself (data_loader.py)."
    )


def slide_how_it_works(prs):
    slide = content_slide(prs, 3, "How It Works", "From a job description to a ranked list, in 5 steps")

    steps = (
        ("Read", "Take the text out of each resume file (PDF, TXT or MD). The job description is typed or pasted."),
        ("Clean", "Make all letters small, remove links, e-mails and extra symbols, fix broken characters. "
                  "Keep special skill names like C++, C# and .NET."),
        ("Convert", "A computer cannot compare words directly, so each resume and the job are turned into "
                    "a list of numbers, called a vector."),
        ("Compare", "Cosine similarity checks how close two vectors point. Same direction = very similar = "
                    "high score."),
        ("Rank", "Sort the resumes from highest to lowest score. Show the top ones and let the user "
                 "download the list as a CSV file."),
    )
    w, gap = 2.2, 0.3
    for i, (title, body) in enumerate(steps):
        x = 0.6 + i * (w + gap)
        card(slide, x, 1.55, w, 3.35, fill=PANEL)
        badge(slide, x + 0.2, 1.75, 0.55, str(i + 1), size=18)
        text(slide, x + 0.2, 2.45, w - 0.4, 0.4, title, size=17, bold=True)
        text(slide, x + 0.2, 2.9, w - 0.35, 1.9, body, size=13, color=ICE)
        if i < len(steps) - 1:
            arrow(slide, x + w + 0.05, 3.08, 0.2, 0.28)

    card(slide, 0.6, 5.15, 12.1, 0.95, fill=PANEL_2)
    text(slide, 0.85, 5.27, 11.6, 0.75,
         [(("What does the score mean?  ", {"bold": True, "color": CYAN}),
           ("100 means the resume and the job point in exactly the same direction (very similar text). 0 means "
            "they have nothing in common. It only measures how alike the texts are. It is not the chance of "
            "getting hired.", {}))],
         size=14, color=ICE)

    pill(slide, 2.1, 6.3, 9.13, 0.5,
         "No chatbot or LLM writes anything here. The app only measures similarity and sorts.",
         fill=NAVY, color=CYAN, border=CYAN, size=13, bold=True)

    NOTES.append(
        "This is how the app works, in five steps.\n\n"
        "Step 1, Read: we take the text out of each resume file. PDF, TXT and MD files are supported.\n"
        "Step 2, Clean: we make everything lowercase, remove links, e-mail addresses and extra symbols, and fix "
        "broken characters. We are careful to keep skill names like C++ and C#, which normal cleaning would "
        "destroy. The job description is cleaned in exactly the same way, so both sides are fair.\n"
        "Step 3, Convert: a computer cannot compare sentences directly, so we turn each text into a list "
        "of numbers called a vector. The four methods on the next slide are four different ways of doing this "
        "step.\n"
        "Step 4, Compare: we use cosine similarity. Imagine each vector as an arrow. If the job's arrow and the "
        "resume's arrow point the same way, the texts are similar and the score is high.\n"
        "Step 5, Rank: we sort the resumes by score and show the top ones.\n\n"
        "Important: the score is only text similarity, not a hiring probability. And there is no chatbot or LLM "
        "generating text; the app only measures and sorts."
    )


def slide_methods(prs):
    slide = content_slide(prs, 4, "Methodology · The 4 Matching Methods", "Four ways to compare a resume with a job")

    methods = (
        ("Matches the same words",
         "Counts the words in each resume. Rare, specific words (like \"Selenium\") get more weight than "
         "common words (like \"work\").",
         "Job asks for \"Spring Boot\" → resumes that mention Spring Boot go up.",
         "Very fast and easy to explain",
         "\"Developer\" and \"programmer\" look like different words to it"),
        ("Knows related words",
         "Uses a ready-made list of 400,000 English words, each stored as 100 numbers. Similar words get similar "
         "numbers. A resume = the average of its words.",
         "\"Doctor\" and \"physician\" are close to each other.",
         "Understands similar words",
         "Averaging blurs details, so most resumes get close scores"),
        ("Understands the meaning",
         "A small pre-trained AI model (bge-small) reads the whole text and turns its meaning into 384 numbers.",
         "\"Built ML models in Python\" matches a Machine Learning job, even with different words.",
         "Understands context and wording",
         "Slower, and reads only about the first page of a resume"),
        ("All three together",
         "Takes the three scores and averages them, like three judges giving marks and taking the average.",
         "To rank high, a resume needs the right skills, related words and the right overall meaning.",
         "Best result in our test; default in the app",
         "Needs all three methods ready (GloVe file + AI model)"),
    )
    w, gap = 2.9, 0.167
    for i, (tagline, simple, example, good, weak) in enumerate(methods):
        x, color = 0.6 + i * (w + gap), METHOD_COLORS[i]
        card(slide, x, 1.5, w, 5.25, fill=PANEL, border=color)
        badge(slide, x + 0.2, 1.68, 0.5, str(i + 1), fill=color, size=16)
        text(slide, x + 0.82, 1.66, w - 0.95, 0.35, METHOD_NAMES[i], size=17, bold=True)
        text(slide, x + 0.82, 1.98, w - 0.95, 0.3, tagline, size=12, color=color, bold=True)
        text(slide, x + 0.2, 2.5, w - 0.4, 0.25, "IN SIMPLE WORDS", size=10, color=MUTED, bold=True)
        text(slide, x + 0.2, 2.74, w - 0.35, 1.3, simple, size=13, color=WHITE)
        text(slide, x + 0.2, 4.1, w - 0.4, 0.25, "EXAMPLE", size=10, color=MUTED, bold=True)
        text(slide, x + 0.2, 4.34, w - 0.35, 0.95, example, size=12, color=ICE, italic=True)
        text(slide, x + 0.2, 5.35, w - 0.35, 0.6, [(("Good: ", {"bold": True, "color": GREEN}), (good, {}))],
             size=12, color=ICE)
        text(slide, x + 0.2, 5.95, w - 0.35, 0.7, [(("Weak: ", {"bold": True, "color": MUTED}), (weak, {}))],
             size=12, color=ICE)

    NOTES.append(
        "Step 3 can be done in different ways. We built four, and all four use the same cosine-similarity "
        "comparison and the same ranking code.\n\n"
        "1. TF-IDF matches the same words. It counts the words in each resume, but rare, specific words like "
        "'Selenium' count more than common words like 'work'. It is very fast and easy to explain, but it only "
        "sees exact words: 'developer' and 'programmer' look unrelated to it.\n\n"
        "2. GloVe knows related words. It is a ready-made list of 400,000 English words, where each word is 100 "
        "numbers, and words with similar meanings have similar numbers. We represent a resume as the average of "
        "its words. The weakness is that averaging a whole resume blurs the details, so most resumes end up with "
        "close scores.\n\n"
        "3. Sentence-BERT understands meaning. A small pre-trained AI model called bge-small reads the whole text "
        "and gives 384 numbers that capture its meaning. So 'built ML models in Python' still matches a Machine "
        "Learning job. It is slower and reads only about the first page (512 word pieces) of each resume. We "
        "first used a model called MiniLM; bge-small scored better in our tests, so we switched.\n\n"
        "4. Hybrid simply averages the three scores, like three judges. A resume needs the right skills, related "
        "words and the right meaning to rank high. This gave the best result, so it is the default in our app."
    )


def slide_implementation(prs):
    slide = content_slide(prs, 5, "Implementation · Live Demo", "Our code and the working app")

    modules = (
        ("resume_matcher/preprocessing.py", "Cleans the text (step 2)"),
        ("resume_matcher/data_loader.py", "Reads resumes, removes empty and duplicate ones"),
        ("resume_matcher/matcher.py", "The 4 matching methods and the ranking (steps 3–5)"),
        ("app.py", "The web page, built with Streamlit"),
        ("evaluate.py", "Tests which method works best on real data"),
    )
    for i, (path, role) in enumerate(modules):
        y = 1.5 + i * 0.97
        card(slide, 0.6, y, 4.75, 0.82, fill=PANEL)
        text(slide, 0.8, y + 0.1, 4.4, 0.3, path, size=13, bold=True, color=CYAN, font=MONO)
        text(slide, 0.8, y + 0.42, 4.4, 0.35, role, size=13, color=ICE)

    shot = ASSETS / "app_screenshot.png"
    with Image.open(shot) as im:
        ratio = im.height / im.width
    shot_w = 6.6
    shot_h = shot_w * ratio
    frame = card(slide, 5.97, 1.47, shot_w + 0.06, shot_h + 0.06, fill=BORDER, border=None, radius=0.02)
    frame.name = "Screenshot frame"
    pic = slide.shapes.add_picture(str(shot), Inches(6.0), Inches(1.5), Inches(shot_w))
    pic.name = "Actual app screenshot"
    text(slide, 6.0, 1.55 + shot_h + 0.06, shot_w, 0.3,
         "Real screenshot: 7 demo resumes, Data Scientist job, Hybrid. Data-science resumes rank first, HR last.",
         size=11, color=MUTED, italic=True)

    steps = ("Upload resumes", "Paste the job", "Pick a method", "Click Match", "See ranking, download CSV")
    w, gap = 1.2, 0.15
    for i, step in enumerate(steps):
        x = 6.0 + i * (w + gap)
        card(slide, x, 6.12, w, 0.72, fill=PANEL_2)
        badge(slide, x + 0.44, 5.96, 0.32, str(i + 1), size=11)
        text(slide, x + 0.06, 6.3, w - 0.12, 0.5, step, size=12, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    NOTES.append(
        "Our code is split into small modules, and each one matches a box in the design. preprocessing.py does "
        "the cleaning. data_loader.py reads resume files and the datasets and removes empty and duplicate "
        "resumes. matcher.py has the four methods and the ranking. app.py is the web page, made with Streamlit. "
        "evaluate.py compares the methods on the real datasets.\n\n"
        "Live demo (run: streamlit run app.py, then open localhost:8501). Upload the seven sample resumes from "
        "demo_samples/pdf, paste the Data Scientist job from demo_samples/job_description.txt, keep the Hybrid "
        "method, and click Match Resumes. The three data-science resumes come first and the HR resume comes last, "
        "which is what we expect. Then show the download button for the CSV. If time allows, switch to TF-IDF "
        "and show that the order changes a little."
    )


def slide_results(prs):
    slide = content_slide(prs, 6, "Evaluation · Testing · Results", "Which method works best?")
    results = load_results()
    tech, mixed = "UpdatedResumeDataSet.csv", "Resume.csv"
    precision = [float(results[(tech, m)]["mean_precision_at_k"]) for m in METHOD_KEYS]
    top1 = [int(results[(tech, m)]["top1_relevant"]) for m in METHOD_KEYS]
    scored = int(results[(tech, METHOD_KEYS[0])]["labelled_jobs"])
    mixed_precision = [float(results[(mixed, m)]["mean_precision_at_k"]) for m in METHOD_KEYS]

    data = CategoryChartData()
    data.categories = METHOD_NAMES
    data.add_series(f"Tech jobs dataset (166 resumes, {scored} jobs)", precision)
    gframe = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.45), Inches(7.3),
                                    Inches(3.9), data)
    gframe.name = "Precision@5 chart"
    chart = gframe.chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "How many of the top 5 resumes were a good fit? (Precision@5)"
    tp = chart.chart_title.text_frame.paragraphs[0].runs[0].font
    tp.size, tp.bold, tp.color.rgb, tp.name = Pt(15), True, WHITE, FONT
    chart.has_legend = False
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.vary_by_categories = False
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format, labels.number_format_is_linked = "0.00", False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size, labels.font.bold, labels.font.color.rgb, labels.font.name = Pt(14), True, WHITE, FONT
    series = plot.series[0]
    for idx, color in enumerate(METHOD_COLORS):
        point = series.points[idx]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = color
    va = chart.value_axis
    va.maximum_scale, va.minimum_scale, va.major_unit = 1.0, 0.0, 0.25
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = BORDER
    va.format.line.fill.background()
    va.tick_labels.font.size, va.tick_labels.font.color.rgb, va.tick_labels.font.name = Pt(12), MUTED, FONT
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = "0.00", False
    ca = chart.category_axis
    ca.tick_labels.font.size, ca.tick_labels.font.color.rgb, ca.tick_labels.font.name = Pt(13), ICE, FONT
    ca.format.line.color.rgb = BORDER

    card(slide, 0.6, 5.45, 7.3, 0.95, fill=PANEL_2)
    text(slide, 0.8, 5.53, 6.95, 0.85,
         [(("How we measured: ", {"bold": True, "color": CYAN}),
           (f"for {scored} tech jobs we looked at the top 5 resumes and checked if their job type fits "
            f"(e.g. Backend job → Java / Python Developer). Hybrid {precision[-1]:.2f} means about "
            f"{round(precision[-1] * 5)} out of 5 were a good fit.", {}))],
         size=13, color=ICE)

    text(slide, 8.3, 1.5, 4.4, 0.3, f"WAS THE #1 RESUME A GOOD FIT?  (out of {scored} jobs)", size=12,
         color=MUTED, bold=True)
    for i, (name, hits, color) in enumerate(zip(METHOD_NAMES, top1, METHOD_COLORS)):
        y = 1.92 + i * 0.45
        text(slide, 8.3, y, 1.45, 0.3, name, size=13, color=WHITE)
        bar(slide, 9.8, y + 0.07, 2.2, 0.17, PANEL_2)
        bar(slide, 9.8, y + 0.07, 2.2 * hits / scored, 0.17, color)
        text(slide, 12.1, y, 0.6, 0.3, f"{hits}/{scored}", size=13, bold=True, align=PP_ALIGN.RIGHT)

    text(slide, 8.3, 3.9, 4.4, 0.3, "TESTING", size=12, color=MUTED, bold=True)
    tiles = ((str(TEST_COUNT), "automated tests, all passing", GREEN),
             ("16", "tech job descriptions written for testing", CYAN))
    for i, (big, label, color) in enumerate(tiles):
        x = 8.3 + i * 2.25
        card(slide, x, 4.28, 2.1, 1.25, fill=PANEL)
        text(slide, x + 0.18, 4.36, 1.8, 0.55, big, size=26, bold=True, color=color)
        text(slide, x + 0.18, 4.9, 1.8, 0.6, label, size=12, color=ICE)
    text(slide, 8.3, 5.7, 4.4, 0.75,
         f"On the all-industries dataset every method scored low ({min(mixed_precision):.2f}–"
         f"{max(mixed_precision):.2f}), because it has very few software resumes.",
         size=12, color=ICE)

    text(slide, 0.6, 6.5, 12.1, 0.4,
         "Honest note: this is a small test, and \"job type fits\" is only a rough check. "
         "The results are a good sign, not a guarantee.",
         size=12, color=MUTED, italic=True)

    NOTES.append(
        "How do we know which method works best? We wrote 16 tech job descriptions, like Backend Developer, "
        "Data Scientist and DevOps. For each job, we let every method pick its top 5 resumes from the tech "
        "dataset, and we checked how many of those 5 come from a fitting job type. For example, for a Backend job, "
        "a Java or Python Developer resume counts as a good fit. This is called Precision at 5. Two jobs, Mobile "
        "and Systems, have no fitting job type in the data, so 14 jobs are scored.\n\n"
        f"Results: Hybrid {precision[3]:.2f}, Sentence-BERT {precision[2]:.2f}, TF-IDF {precision[0]:.2f}, "
        f"GloVe {precision[1]:.2f}. So with Hybrid, about {round(precision[3] * 5)} out of 5 top resumes were a "
        f"good fit, and the number-1 resume was a good fit for {top1[3]} of the {scored} jobs. TF-IDF does well "
        "because tech jobs use exact skill names like Spring Boot or Selenium. GloVe does worst because averaging "
        "blurs these specific words.\n\n"
        f"On the all-industries dataset all methods scored low, around {min(mixed_precision):.2f} to "
        f"{max(mixed_precision):.2f}, because it has very few "
        "software resumes.\n\n"
        f"We also have {TEST_COUNT} automated tests (pytest) that check loading, cleaning, every method, ranking "
        "and the web page. All pass.\n\n"
        "To be honest about limits: 14 jobs is a small test, and 'job type fits' is only a rough check of "
        "relevance, so these numbers are a good sign, not a guarantee. All numbers come from evaluate.py "
        "(results/evaluation_summary.csv)."
    )


def slide_conclusion(prs):
    slide = content_slide(prs, 7, "Milestone 3 · Limitations · Next Steps", "Summary")

    criteria = (
        ("Design and code match",
         ("Each box in our design is a real module in the code",
          "Resumes and jobs are cleaned in exactly the same way",
          "All 4 methods use the same compare-and-rank code")),
        ("Working modules and demo",
         ("4 matching methods tested on 2 real datasets",
          "Web app: upload resumes, rank, download CSV",
          f"Evaluation script and {TEST_COUNT} automated tests")),
    )
    for i, (title, items) in enumerate(criteria):
        y = 1.5 + i * 2.25
        card(slide, 0.6, y, 6.4, 2.05)
        badge(slide, 0.82, y + 0.22, 0.5, "✓", fill=GREEN, size=18)
        text(slide, 1.5, y + 0.25, 5.3, 0.4, title, size=17, bold=True)
        bullets(slide, 1.5, y + 0.8, 5.3, 1.2, items, size=14, bullet_color=GREEN, space_after=4)

    for i, (title, color, items) in enumerate((
        ("Limitations", MUTED, ("Small test: only 16 job descriptions",
                                "\"Job type fits\" is only a rough check",
                                "No Word (DOCX) files or scanned PDFs yet")),
        ("Next steps", CYAN, ("Ask real people to mark good matches",
                              "Show which skills match or are missing",
                              "Support DOCX files and scanned PDFs")),
    )):
        y = 1.5 + i * 2.25
        card(slide, 7.3, y, 5.4, 2.05, fill=PANEL)
        text(slide, 7.55, y + 0.22, 4.9, 0.4, title, size=17, bold=True, color=color)
        bullets(slide, 7.55, y + 0.75, 4.95, 1.25, items, size=14, bullet_color=color, space_after=4)

    text(slide, 0.6, 6.1, 4, 0.7, "Thank You", size=34, bold=True)
    # Plain text, not a hyperlink: Keynote draws links in dark blue, which is unreadable on the navy background.
    text(slide, 4.6, 6.28, 8.1, 0.4, REPO_URL, size=15, color=CYAN, align=PP_ALIGN.RIGHT)

    NOTES.append(
        "To sum up the two Milestone 3 points. First, design and code match: every box in our design is a real "
        "module, resumes and jobs are cleaned the same way, and all four methods share the same compare-and-rank "
        "code. Second, the modules work: four methods tested on two real datasets, a working web app, an "
        f"evaluation script and {TEST_COUNT} automated tests.\n\n"
        "Our limitations: the test is small, 'job type fits' is only a rough check, and we cannot read Word files "
        "or scanned PDFs yet. Next, we want real people to mark good matches, show which skills match or are "
        "missing, and support more file types.\n\n"
        "Thank you. The code is on GitHub at the link shown. We are happy to take questions."
    )


# Simple answers to questions the panel may ask (written to speaker_notes.md only).
QA = (
    ("What is cosine similarity?",
     "Think of each text as an arrow. Cosine similarity measures the angle between the job's arrow and the "
     "resume's arrow. Same direction gives 1 (we show 100), at right angles gives 0. It ignores length, so a long "
     "resume is not favoured just for being long."),
    ("What does TF-IDF stand for?",
     "Term Frequency, Inverse Document Frequency. Term frequency: how often a word appears in this resume. "
     "Inverse document frequency: words that appear in almost every resume (like 'work') get a low weight, rare "
     "words (like 'Kubernetes') get a high weight."),
    ("Is this an LLM or ChatGPT-style app?",
     "No. Nothing generates text. Sentence-BERT is a small pre-trained model that only turns text into numbers. "
     "The app measures similarity and sorts."),
    ("Why is GloVe so weak here?",
     "We average all the word vectors in a resume. Most tech resumes share many general computer words, so their "
     "averages look alike and the specific skills get blurred."),
    ("Why use the Hybrid?",
     "Each method catches something different: exact skills (TF-IDF), related words (GloVe) and overall meaning "
     "(Sentence-BERT). Averaging them gave the best result on the tech dataset and put a good fit at number 1 "
     "for all 14 scored jobs."),
    ("Why did you switch the Sentence-BERT model?",
     "We first used all-MiniLM-L6-v2 (Precision@5 0.69). BAAI/bge-small-en-v1.5 is about the same size, is "
     "trained for search, and scored 0.81 on the same test, so it became the default."),
    ("Is the score the chance of being hired?",
     "No. It is only how similar the texts are. The app should help a recruiter decide what to read first, never "
     "reject anyone automatically. Text similarity can also copy biases in how people write."),
    ("Is the data real?",
     "Yes. Both datasets are public Kaggle resume datasets, and all results come from real runs. The 7 demo resumes in "
     "demo_samples/ are made-up sample resumes prepared only for the live demo, so we show no real person's data."),
)


def write_notes() -> None:
    lines = ["# Speaker notes — Milestone 3", "",
             "Simple script for each slide. Say it in your own words; you do not need to read it exactly.", ""]
    for number, note in enumerate(NOTES, start=1):
        lines += [f"## Slide {number}", "", note, ""]
    lines += ["## Possible questions (simple answers)", ""]
    for question, answer in QA:
        lines += [f"**{question}**", "", answer, ""]
    NOTES_OUTPUT.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    ensure_assets()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = "Resume Matcher — Milestone 3"
    prs.core_properties.author = "Yash Agarwal, Raghav Soni (Batch 3)"
    for build in (slide_title, slide_problem, slide_how_it_works, slide_methods, slide_implementation,
                  slide_results, slide_conclusion):
        build(prs)
    prs.save(OUTPUT)
    print(f"Saved {OUTPUT} ({len(prs.slides)} slides)")
    write_notes()
    print(f"Saved {NOTES_OUTPUT}")


if __name__ == "__main__":
    main()
